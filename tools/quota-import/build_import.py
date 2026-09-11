# -*- coding: utf-8 -*-
"""EK 1. Annex 1 Quotas -> orhun_quotas import tayyorlash."""
import openpyxl, re, json, csv

XL = "/home/tiu/Downloads/Telegram Desktop/EK 1. Annex 1 Quotas (1).xlsx"
wb = openpyxl.load_workbook(XL, data_only=True)

# ---- 1. Nom moslashuvi: 74 tasi avtomatik, 13 tasini qo'lda aniqladim -------
MANUAL = {   # Excel qatori -> member id
    5: 127,   # Azerbaycan Devlet Pedagoji Ü.  -> Azerbaijan State Pedagogical University
    21: 38,   # Karaganda ... Buketov          -> Karaganda Buketov University
    25: 28,   # Hoca Ahmet Yesevi Türk-Kazak   -> Akhmet Yassawi University
    26: 30,   # Kh.Dosmukhamedov Atyrau        -> Atyrau State University
    35: 45,   # Jusup Balasagyn Kyrgyz Nat.    -> Kyrgyz National University
    36: 46,   # Bishkek State U. (K.Karasayev) -> Bishkek Humanities University  (*)
    37: 47,   # Int. U. of the Kyrgyz Republic -> International University of Kyrgyzstan (*)
    43: 117,  # Özbekistan Gazetecilik         -> Journalism and Mass Communications University
    44: 120,  # Özbekistan devlet sanat        -> Uzbekistan State Institute of Arts and Culture
    48: 121,  # Termez                         -> Termiz University of Economics and Service
    60: 63,   # Kahramanmaraş Sütçü İmam       -> Kahramanmaraş Sütçüimam University
    64: 150,  # Adana Alparslan Türkeş         -> Adana Alparslan Türkeş Science and Technology U.
    81: 85,   # Ankara Müzik ve Güzel Sanatlar -> Ankara Müzik ve Güzel Sanatlar Üniversitesi
}

# ---- 2. Raqamlarni o'qish ---------------------------------------------------
NONE_WORDS = {"yok", "x", "-", "--", "---", "yoktur", "sağlanamayacaktır", "none", "no"}

def parse_num(v):
    """Excel katagidan bir o'quv yiliga mo'ljallangan JAMI kvotani chiqaradi."""
    if v is None: return None, ""
    s = re.sub(r"\s+", " ", str(v)).strip()
    if not s: return None, ""
    if s.lower() in NONE_WORDS: return 0, ""
    # "Güz/ Fall (8) Bahar/ Spring (7) Toplam / Total (15)"
    m = re.search(r"(?:toplam|total)\s*[/\w]*\s*\((\d+)\)", s, re.I)
    if m: return int(m.group(1)), s
    # "10/10 (20)" yoki "5/5(10)" -> qavs ichidagi jami
    m = re.fullmatch(r"(\d+)\s*/\s*(\d+)\s*\(\s*(\d+)\s*\)?", s)
    if m:
        a, b, tot = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return tot, f"Fall/Spring {a}/{b}"
    # "10/20", "20/20"
    m = re.fullmatch(r"(\d+)\s*/\s*(\d+)", s)
    if m:
        a, b = int(m.group(1)), int(m.group(2))
        return a + b, f"Fall/Spring {a}/{b}"
    # "Güz 4 Bahar 4"  /  "5 bahar 5 güz"
    guz = re.search(r"(?:güz|guz|fall)\D{0,3}(\d+)|(\d+)\s*(?:güz|guz|fall)", s, re.I)
    bah = re.search(r"(?:bahar|spring)\D{0,3}(\d+)|(\d+)\s*(?:bahar|spring)", s, re.I)
    if guz and bah:
        g = int(guz.group(1) or guz.group(2)); b = int(bah.group(1) or bah.group(2))
        return g + b, f"Fall/Spring {g}/{b}"
    if guz and not bah:
        g = int(guz.group(1) or guz.group(2)); return g, "Fall term only"
    if bah and not guz:
        b = int(bah.group(1) or bah.group(2)); return b, "Spring term only"
    # "3 (spring)" / "5 (fall)" -> bitta semestrga tegishli son
    m = re.fullmatch(r"(\d+)\s*\(\s*(güz|guz|fall|bahar|spring)\s*\)", s, re.I)
    if m:
        kind = m.group(2).lower()
        return int(m.group(1)), "Fall term only" if kind in ("güz","guz","fall") else "Spring term only"
    # "12*  (Bkz. Notlar -M sütunu)" -> yulduzcha izohga havola, soni 12
    m = re.match(r"^(\d+)\s*\*", s)
    if m: return int(m.group(1)), ""
    # "1 akademik yıl için 8"  -> oxirgi son
    if re.search(r"akademik|academic", s, re.I):
        nums = re.findall(r"\d+", s)
        if nums: return int(nums[-1]), s
    # "12*", "10.0", "31"
    m = re.match(r"^(\d+)(?:\.0+)?\s*\*?", s)
    if m and not re.search(r"[a-zA-ZçğıöşüÇĞİÖŞÜ]{4,}", s[m.end():m.end()+12]):
        return int(m.group(1)), ""
    return None, s          # tushunarsiz -> qo'lda ko'rish uchun qaytariladi

# ---- 2b. Qo'lda tuzatish: Excel katagining o'zi ziddiyatli bo'lgan joylar ---
# "2/2 (40)" kabi kataklarda qavs ichidagi son semestrlar yig'indisiga to'g'ri kelmaydi.
# User qarori (2026-09-11): qavsdagi kattasi emas, semestrlar yig'indisi olinsin.
OVERRIDE = {
    (41, "sa"): (4,  "Excel: \"2/2 (40)\" — 2+2=4 olindi"),   # Tashkent International University
    (45, "sa"): (10, "Excel: \"5/5 (40)\" — 5+5=10 olindi"),  # Samarkand State University
}

# ---- 3. Batafsil varaqlar (asosiy jadvalda "see Nth sheet" deb yozilganlar) --
# (varaq nomi, boshlanish qatori, {maydon: ustun indeksi})
DETAIL = {
  51: ("Erzincan Binali Yıldırım",        6, {"sa": 4, "ta": 5}),                      # faqat "kabul edeceği"
  57: ("Selçuk Üniversitesi",             6, {"ss": 4, "sa": 5, "ts": 6, "ta": 7}),
  59: ("Anadolu Üniversitesi",            7, {"ss": 5, "sa": 6, "ts": 7, "ta": 8}),
  60: ("Kahramanmaraş Sütçü İmam Üniver", 7, {"ss": 4, "sa": 5, "ts": 6, "ta": 7}),
  61: ("Ege University",                  5, {"sa": 4, "ta": 5, "ss": 6, "ts": 7}),    # ustun tartibi boshqacha!
  73: ("KTO Karatay Üniversitesi",        6, {"ss": 4, "sa": 5, "ts": 6, "ta": 7}),
  78: ("AAKÜ",                            6, {"ss": 4, "sa": 5, "ts": 6, "ta": 7}),
}

def sum_detail(sheet, start, cols):
    ws = wb[sheet]; tot = {k: 0 for k in cols}; n = 0
    for r in ws.iter_rows(min_row=start, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
        vals = {}
        for k, ci in cols.items():
            v, _ = parse_num(r[ci] if ci < len(r) else None)
            vals[k] = v
        if any(v is not None for v in vals.values()):
            n += 1
            for k, v in vals.items(): tot[k] += (v or 0)
    return tot, n

# ---- 4. Izoh matni (512 belgi) ---------------------------------------------
def clean(v, limit=None):
    if v is None: return ""
    s = re.sub(r"\s+", " ", str(v)).strip(" .-*")
    return s[:limit] if limit else s

def english_half(s):
    """'Var/Available' -> 'Available';  'EVET' -> 'Yes'."""
    s = clean(s)
    if not s: return ""
    if "/" in s:
        parts = [p.strip() for p in s.split("/") if p.strip()]
        ascii_parts = [p for p in parts if all(ord(c) < 128 for c in p)]
        if ascii_parts: s = max(ascii_parts, key=len)
    tr = {"var": "Yes", "evet": "Yes", "yok": "No", "yoktur": "No", "hayir": "No",
          "sağlanamayacaktır": "Not available", "not provided": "Not provided"}
    return tr.get(s.lower(), s)

def term_text(v):
    """'Güz/Bahar Dönemi Fall/Spring Term' -> 'Fall and Spring' (bo'lakka bo'lib yuborilmasin)."""
    s = clean(v)
    if not s: return ""
    low = s.lower()
    fall   = any(w in low for w in ("güz", "guz", "fall"))
    spring = any(w in low for w in ("bahar", "spring"))
    if fall and spring: return "Fall and Spring"
    if fall:   return "Fall"
    if spring: return "Spring"
    return s[:70]

LBL = {"ss": "students out", "sa": "students in", "ts": "staff out", "ta": "staff in"}

def build_note(r, splits):
    bits = []
    if "_" in splits:
        bits.append(splits.pop("_"))
    uniq = set(splits.values())
    if len(uniq) == 1:
        bits.append(uniq.pop())                       # hamma ustunda bir xil taqsimot
    elif uniq:
        bits.append("; ".join(f"{LBL[k]} {v}" for k, v in sorted(splits.items())))
    term = term_text(r[10])
    if term: bits.append(f"Term: {term}")
    for label, ci, lim in (("Dormitory", 7, 120), ("Meals", 8, 100)):
        c = english_half(r[ci])[:lim]
        if c: bits.append(f"{label}: {c}")
    note = " · ".join(bits)
    return note[:509] + "..." if len(note) > 512 else note

# ---- 5. Asosiy jadvalni qayta ishlash ---------------------------------------
auto = {o["row"]: o["cands"][0] for o in json.load(open("cands.json"))
        if o["cands"][0]["score"] >= 0.80 and o["cands"][0]["score"] - o["cands"][1]["score"] >= 0.12}
meta = {o["row"]: o for o in json.load(open("cands.json"))}
members = {int(l.split("\t")[0]): l.rstrip("\n").split("\t")[2] for l in open("members.tsv", encoding="utf-8")}

ws = wb["Genel Tablo"]
out, problems = [], []
for i, r in enumerate(ws.iter_rows(min_row=3, max_row=90, max_col=13, values_only=True), start=3):
    if not r[2] or not str(r[2]).strip(): continue
    mid = MANUAL.get(i) or (auto[i]["id"] if i in auto else None)
    if mid is None: problems.append((i, "member topilmadi", meta[i]["uni"])); continue

    vals, splits, src = {}, {}, "asosiy jadval"
    if i in DETAIL:
        sheet, start, cols = DETAIL[i]
        tot, n = sum_detail(sheet, start, cols)
        vals = {"ss": tot.get("ss"), "sa": tot.get("sa"), "ts": tot.get("ts"), "ta": tot.get("ta")}
        src = f'"{sheet}" varag\'i, {n} qator yig\'indisi'
        splits["_"] = f"Faculty-level breakdown: {n} programmes"
    else:
        for key, ci in (("ss", 3), ("sa", 4), ("ts", 5), ("ta", 6)):
            v, hint = parse_num(r[ci])
            if (i, key) in OVERRIDE:
                v, _why = OVERRIDE[(i, key)]
            vals[key] = v
            if hint and not hint.startswith(("Fall","Spring")) and v is None:
                problems.append((i, f"raqam o'qilmadi ({key}): {hint[:60]}", meta[i]["uni"]))
            if hint.startswith(("Fall","Spring")): splits[key] = hint

    out.append({
        "row": i, "member_id": mid, "member": members[mid], "xl_name": meta[i]["uni"],
        "country": meta[i]["country"], "source": src,
        "student_send": vals["ss"], "student_accept": vals["sa"],
        "teacher_send": vals["ts"], "teacher_accept": vals["ta"],
        "note": build_note(r, splits),
        "skipped_departments": clean(r[9])[:60],
    })

json.dump(out, open("import_rows.json", "w"), ensure_ascii=False, indent=1)
with open("import_preview.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(out[0].keys())); w.writeheader(); w.writerows(out)

print(f"Tayyor qatorlar: {len(out)}")
print("Muammoli:", len(problems))
for p in problems: print("   ", p)
ids = {}
for o in out: ids.setdefault(o["member_id"], []).append(o["row"])
print("Dublikat member:", {members[k]: v for k, v in ids.items() if len(v) > 1} or "yo'q")
miss = [o["row"] for o in out if all(o[k] is None for k in ("student_send","student_accept","teacher_send","teacher_accept"))]
print("To'rt raqami ham bo'sh qatorlar:", miss or "yo'q")
